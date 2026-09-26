import cv2
import numpy as np
import os
import logging
from datetime import datetime
from scipy.signal import savgol_filter
from scipy.stats import linregress
import matplotlib.pyplot as plt
import shutil
from openpyxl import Workbook

# ====== Configuration ======
CANNY_THRESHOLD = 8
MEDIAN_BLUR_K_SIZE = 9
MORPH_K_SIZE = 13
FRAME_RATE = 30.0
DISPLAY_RESOLUTION = (640, 480)
ROI_EXPAND = 80

kernel = np.ones((MORPH_K_SIZE, MORPH_K_SIZE), np.uint8)

# ====== Paths ======
INPUT_FOLDERS = [
    r"C:\Users\madhu\Desktop\Pupil project\Assets\frames4"
]

BASE_OUTPUT = r"C:\Users\madhu\Desktop\Pupil project\Outputs"


def create_workspace(input_folder_name):
    timestamp = datetime.now().strftime("%d_%m_%Y_%H_%M_%S")
    workspace_name = f"{input_folder_name}_Ellipse_{timestamp}"
    workspace = os.path.join(BASE_OUTPUT, workspace_name)
    os.makedirs(workspace, exist_ok=True)

    zero_area_img_folder = os.path.join(workspace, "zero_area_frames")
    area_detected_img_folder = os.path.join(workspace, "area_detected_frames")
    ellipse_fitted_folder = os.path.join(workspace, "ellipse_fitted_frames")
    yolo_label_folder = os.path.join(workspace, "YOLO data")

    os.makedirs(zero_area_img_folder, exist_ok=True)
    os.makedirs(area_detected_img_folder, exist_ok=True)
    os.makedirs(ellipse_fitted_folder, exist_ok=True)
    os.makedirs(yolo_label_folder, exist_ok=True)

    return workspace, zero_area_img_folder, area_detected_img_folder, ellipse_fitted_folder, yolo_label_folder


def setup_logger(folder_name):
    logger = logging.getLogger(f"EllipseLogger_{os.path.basename(folder_name)}")
    logger.setLevel(logging.DEBUG)

    if logger.hasHandlers():
        logger.handlers.clear()

    handler = logging.FileHandler(os.path.join(folder_name, "log_ellipse_fitting.log"))
    handler.setFormatter(logging.Formatter('%(asctime)s - %(levelname)s - %(message)s'))
    logger.addHandler(handler)
    return logger


def process_frame(img, roi=None):
    if roi is not None:
        x, y, w, h = roi
        sub_img = img[y:y+h, x:x+w]
        gray = cv2.cvtColor(sub_img, cv2.COLOR_BGR2GRAY)
        gray = cv2.medianBlur(gray, MEDIAN_BLUR_K_SIZE)
        gray = cv2.morphologyEx(gray, cv2.MORPH_OPEN, kernel)
        edges = cv2.Canny(gray, CANNY_THRESHOLD, CANNY_THRESHOLD * 2)

        mask = np.zeros(img.shape[:2], dtype=np.uint8)
        mask[y:y+h, x:x+w] = edges
        return mask
    else:
        gray = cv2.cvtColor(img, cv2.COLOR_BGR2GRAY)
        gray = cv2.medianBlur(gray, MEDIAN_BLUR_K_SIZE)
        gray = cv2.morphologyEx(gray, cv2.MORPH_OPEN, kernel)
        return cv2.Canny(gray, CANNY_THRESHOLD, CANNY_THRESHOLD * 2)


def filter_contour(contours):
    filtered = []
    for c in contours:
        try:
            hull = cv2.convexHull(c)
            area = cv2.contourArea(hull)
            perimeter = cv2.arcLength(hull, True)
            circularity = (4 * np.pi * area) / (perimeter ** 2) if perimeter > 0 else 0
            if circularity > 0.98:
                filtered.append(hull)
        except:
            continue
    return filtered


# ===============================
# MAIN PROCESSING FUNCTION
# ===============================
def process_single_folder(input_folder):
    folder_name = os.path.basename(os.path.normpath(input_folder))
    print(f"--- Processing folder: {folder_name} ---")

    workspace, zero_area_img_folder, area_detected_img_folder, ellipse_fitted_folder, yolo_label_folder = create_workspace(folder_name)
    logger = setup_logger(workspace)

    area_txt_path = os.path.join(workspace, "area.txt")

    if not os.path.exists(input_folder):
        print(f"Error: Folder {input_folder} does not exist.")
        return

    images = sorted([f for f in os.listdir(input_folder)
                     if f.lower().endswith((".png", ".jpg", ".jpeg"))])

    if not images:
        print(f"No images found in {input_folder}.")
        return

    frame_idx = 0
    area_list, time_list = [], []
    prev_ellipse = None

    fourcc = cv2.VideoWriter_fourcc(*'mp4v')
    video_writer = cv2.VideoWriter(
        os.path.join(workspace, "recorded.mp4"),
        fourcc,
        FRAME_RATE,
        DISPLAY_RESOLUTION
    )

    zero_wb = Workbook()
    zero_ws = zero_wb.active
    zero_ws.title = "Zero Area Frames"
    zero_ws.append(["Serial No", "Frame No", "Time (s)", "Area"])

    detected_wb = Workbook()
    detected_ws = detected_wb.active
    detected_ws.title = "Detected Area Frames"
    detected_ws.append(["Serial No", "Frame No", "Time (s)", "Area"])

    all_wb = Workbook()
    all_ws = all_wb.active
    all_ws.title = "All Frames"
    all_ws.append(["Serial No", "Frame No", "Time (s)", "Area"])

    zero_serial = detected_serial = all_serial = 1

    with open(area_txt_path, 'w') as f:
        f.write("Area, Frame, Time\n")

    for img_file in images:
        img_path = os.path.join(input_folder, img_file)
        frame = cv2.imread(img_path)
        if frame is None:
            continue

        frame = cv2.resize(frame, DISPLAY_RESOLUTION)
        img_h, img_w = frame.shape[:2]
        elapsed = frame_idx / FRAME_RATE

        roi = None
        if prev_ellipse is not None:
            x_prev, y_prev, MA_prev = prev_ellipse
            size = int(MA_prev / 2) + ROI_EXPAND
            x1 = max(0, int(x_prev - size))
            y1 = max(0, int(y_prev - size))
            x2 = min(img_w, int(x_prev + size))
            y2 = min(img_h, int(y_prev + size))
            roi = (x1, y1, x2 - x1, y2 - y1)

        canny_output = process_frame(frame, roi)
        contours, _ = cv2.findContours(canny_output, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE)
        filtered = filter_contour(contours)

        latest_area = 0.0
        ellipse_detected = False

        for c in filtered:
            if len(c) >= 5:
                ellipse = cv2.fitEllipse(c)
                (x, y), (MA, ma), _ = ellipse
                area = (np.pi / 4) * MA * ma

                cv2.ellipse(frame, ellipse, (0, 255, 0), 2)

                with open(area_txt_path, 'a') as f:
                    f.write(f"{area:.2f}, {frame_idx}, {elapsed:.2f}\n")

                # ===== YOLO LABEL SAVE =====
                class_id = 0
                x_center = x / img_w
                y_center = y / img_h
                box_width = MA / img_w
                box_height = ma / img_h

                label_path = os.path.join(
                    yolo_label_folder,
                    os.path.splitext(img_file)[0] + ".txt"
                )

                with open(label_path, "w") as f:
                    f.write(f"{class_id} {x_center:.6f} {y_center:.6f} {box_width:.6f} {box_height:.6f}\n")

                area_list.append(area)
                time_list.append(elapsed)

                latest_area = area
                prev_ellipse = (x, y, MA)
                ellipse_detected = True

                cv2.imwrite(os.path.join(ellipse_fitted_folder, img_file), frame)
                break

        if not ellipse_detected:
            prev_ellipse = None

        if latest_area == 0.0:
            zero_ws.append([zero_serial, frame_idx, round(elapsed, 2), latest_area])
            shutil.copy(img_path, os.path.join(zero_area_img_folder, img_file))
            zero_serial += 1
        else:
            detected_ws.append([detected_serial, frame_idx, round(elapsed, 2), round(latest_area, 2)])
            shutil.copy(img_path, os.path.join(area_detected_img_folder, img_file))
            detected_serial += 1

        all_ws.append([all_serial, frame_idx, round(elapsed, 2), round(latest_area, 2)])
        all_serial += 1

        overlay_text = f"Frame: {frame_idx} | Time: {elapsed:.2f}s | Area: {latest_area:.1f}"
        cv2.putText(frame, overlay_text, (10, 30),
                    cv2.FONT_HERSHEY_SIMPLEX, 0.7,
                    (0, 255, 255), 2, cv2.LINE_AA)

        video_writer.write(frame)
        cv2.imshow("Ellipse Detection with ROI Tracking + YOLO Labels", frame)

        if cv2.waitKey(1) & 0xFF in [ord('q'), 27]:
            print("User interrupted processing.")
            break

        frame_idx += 1

    zero_wb.save(os.path.join(workspace, "zero_area_frames.xlsx"))
    detected_wb.save(os.path.join(workspace, "area_detected_frames.xlsx"))
    all_wb.save(os.path.join(workspace, "all_frames.xlsx"))

    video_writer.release()
    cv2.destroyAllWindows()

    print(f"--- Finished processing {folder_name} ---\n")


def main():
    for folder in INPUT_FOLDERS:
        process_single_folder(folder)
    print("All folders processed successfully!")


if __name__ == "__main__":
    main()